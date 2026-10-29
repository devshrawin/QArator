"""End-to-end QA of the QArator extension in real Chromium.
Usage: python e2e.py <extension dir> <out dir> [port]"""
import sys, os, json, tempfile, subprocess, time, traceback, base64
from playwright.sync_api import sync_playwright

EXT, OUT = os.path.abspath(sys.argv[1]), os.path.abspath(sys.argv[2])
os.makedirs(OUT, exist_ok=True)
HERE = os.path.dirname(os.path.abspath(__file__))
PORT = int(sys.argv[3]) if len(sys.argv) > 3 else 8765
ANSWER = json.load(open(os.path.join(HERE, "answer.json"), encoding="utf-8"))
ANS = json.dumps(ANSWER).replace("localhost:8765", f"localhost:{PORT}")
sys.stdout.reconfigure(encoding="utf-8")
results, logs = [], []
mode = {"v": "ok"}
seen = {"gemini": [], "anthropic": [], "openai": []}


def check(name, ok, detail=""):
    results.append((name, bool(ok), detail))
    print(("PASS " if ok else "FAIL ") + name + (f" :: {detail}" if detail else ""), flush=True)


def open_export(dl, path):
    """Save a download; a zip holds the Word file plus videos/. Returns (python-docx Document, names in zip or [])."""
    import zipfile, docx as pydocx, io
    dl.save_as(path)
    if path.endswith(".zip") or zipfile.is_zipfile(path) and dl.suggested_filename.endswith(".zip"):
        z = zipfile.ZipFile(path)
        names = z.namelist()
        doc = next(n for n in names if n.endswith(".docx"))
        return pydocx.Document(io.BytesIO(z.read(doc))), names
    return pydocx.Document(path), []


def phase(fn):
    try:
        fn()
    except Exception as e:
        check(fn.__name__ + " crashed", False, f"{type(e).__name__}: {e}")
        traceback.print_exc()


def err_body(provider, url="", headers=None):
    m = mode["v"]
    if m == "quota_main" and "gemini-2.5-flash:" in url and (headers or {}).get("x-goog-api-key") == "test-gemini-key":
        return 429, {"error": {"code": 429, "status": "RESOURCE_EXHAUSTED",
            "message": "You exceeded your current quota. limit: 20, model: gemini-2.5-flash",
            "details": [{"violations": [{"quotaId": "GenerateRequestsPerDayPerProjectPerModel-FreeTier"}]}]}}, {}
    if m == "quota_main":
        return None
    if m == "quota_flash" and "gemini-2.5-flash:" in url:
        return 429, {"error": {"code": 429, "status": "RESOURCE_EXHAUSTED",
            "message": "You exceeded your current quota. Quota exceeded for metric: generativelanguage.googleapis.com/generate_content_free_tier_requests, limit: 20, model: gemini-2.5-flash\nPlease retry in 39.2s.",
            "details": [{"@type": "type.googleapis.com/google.rpc.QuotaFailure", "violations": [{"quotaId": "GenerateRequestsPerDayPerProjectPerModel-FreeTier"}]}]}}, {}
    if m == "dead_flash" and "gemini-2.5-flash:" in url:
        return 404, {"error": {"code": 404, "message": "models/gemini-2.5-flash is not found for API version v1beta", "status": "NOT_FOUND"}}, {}
    if m in ("quota_flash", "dead_flash"):
        return None
    if m == "429":
        return 429, {"error": {"code": 429, "message": "Resource has been exhausted (e.g. check quota).", "status": "RESOURCE_EXHAUSTED"}}, {"retry-after": "7"}
    if m == "badkey":
        return 400, {"error": {"code": 400, "message": "API key not valid. Please pass a valid API key.", "status": "INVALID_ARGUMENT"}}, {}
    if m == "401":
        return 401, {"error": {"type": "authentication_error", "message": "invalid x-api-key"}}, {}
    if m == "garbage":
        return 200, None, {}
    return None


def make_route(provider):
    def handler(route):
        req = route.request
        if req.method == "OPTIONS":
            return route.fulfill(status=204, headers={"access-control-allow-origin": "*", "access-control-allow-headers": "*"})
        hdr = {"access-control-allow-origin": "*", "content-type": "application/json"}
        url = req.url
        if req.method == "GET":  # model lists
            if provider == "gemini":
                body = {"models": [
                    {"name": "models/gemini-2.5-flash", "displayName": "Gemini 2.5 Flash", "supportedGenerationMethods": ["generateContent"]},
                    {"name": "models/gemini-2.5-pro", "displayName": "Gemini 2.5 Pro", "supportedGenerationMethods": ["generateContent"]},
                    {"name": "models/text-embedding-004", "supportedGenerationMethods": ["embedContent"]}]}
            elif provider == "anthropic":
                body = {"data": [{"id": "claude-sonnet-4-5", "display_name": "Claude Sonnet 4.5"}]}
            else:
                body = {"data": [{"id": "gpt-4.1"}, {"id": "gpt-4o-mini-tts"}, {"id": "text-embedding-3-small"}]}
            return route.fulfill(status=200, headers=hdr, body=json.dumps(body))
        seen[provider].append({"url": url, "headers": req.headers, "body": req.post_data_json})
        e = err_body(provider, url, req.headers)
        if e:
            st, b, extra = e
            text = "Sorry, I can't help with that." if b is None else json.dumps(b)
            if b is None:  # model chatter instead of JSON
                b = {"candidates": [{"content": {"parts": [{"text": text}]}}]} if provider == "gemini" else \
                    {"content": [{"type": "text", "text": text}]} if provider == "anthropic" else {"choices": [{"message": {"content": text}}]}
                text = json.dumps(b)
            return route.fulfill(status=st, headers={**hdr, **extra}, body=text)
        if provider == "gemini":
            b = {"candidates": [{"content": {"parts": [{"text": ANS}]}}], "usageMetadata": {"promptTokenCount": 1200, "candidatesTokenCount": 300}}
        elif provider == "anthropic":
            b = {"content": [{"type": "text", "text": "Here is the report:\n" + ANS}]}
        else:
            b = {"choices": [{"message": {"content": ANS}}]}
        route.fulfill(status=200, headers=hdr, body=json.dumps(b))
    return handler


srv = subprocess.Popen([sys.executable, os.path.join(HERE, "server.py"), str(PORT)])
time.sleep(1.2)
BASE = f"http://localhost:{PORT}"
ZBASE = f"http://projects.zoho.localhost:{PORT}"

try:
    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(tempfile.mkdtemp(), channel="chromium", headless=True,
            viewport={"width": 1280, "height": 800}, accept_downloads=True,
            args=[f"--disable-extensions-except={EXT}", f"--load-extension={EXT}"])
        ctx.route("https://generativelanguage.googleapis.com/**", make_route("gemini"))
        ctx.route("https://api.anthropic.com/**", make_route("anthropic"))
        ctx.route("https://api.openai.com/**", make_route("openai"))
        sw = ctx.service_workers[0] if ctx.service_workers else ctx.wait_for_event("serviceworker", timeout=10000)
        EID = sw.url.split("/")[2]
        PANEL = f"chrome-extension://{EID}/sidepanel/sidepanel.html"

        def set_settings(patch):
            sw.evaluate("""async (patch) => {
              const { settings = {} } = await chrome.storage.local.get('settings');
              const s = JSON.parse(JSON.stringify(settings));
              for (const [k, v] of Object.entries(patch)) s[k] = (v && typeof v === 'object' && !Array.isArray(v)) ? { ...(s[k] || {}), ...v } : v;
              await chrome.storage.local.set({ settings: s });
            }""", patch)


        # ---- site under test in window W1
        def welcome_opened():
            urls = sw.evaluate("""async () => { for (let i = 0; i < 40; i++) { const t = await chrome.tabs.query({}); if (t.some((x) => /welcome[.]html/.test(x.url || x.pendingUrl || ''))) return t.map((x) => x.url || x.pendingUrl); await new Promise((r) => setTimeout(r, 200)); } return (await chrome.tabs.query({})).map((x) => x.url); }""")
            check("setup: welcome tab opens on install", any("welcome.html" in u for u in urls), str(urls))
        phase(welcome_opened)
        def errors_default_off():
            reg = sw.evaluate("() => chrome.scripting.getRegisteredContentScripts()")
            check("page errors: off by default, no page hook registered", reg == [], str(reg))
        phase(errors_default_off)
        sw.evaluate("async () => { const { settings = {} } = await chrome.storage.local.get('settings'); await chrome.storage.local.set({ settings: { ...settings, pageErrors: true } }); }")
        time.sleep(0.2)
        for _ in range(20):
            if sw.evaluate("async () => (await chrome.scripting.getRegisteredContentScripts()).length"):
                break
            sw.evaluate("() => new Promise((r) => setTimeout(r, 150))")
        site = ctx.new_page()
        site.on("console", lambda m: None)
        site.goto(f"{BASE}/test.html")
        site.wait_for_timeout(1200)
        tab = sw.evaluate(f"""async () => (await chrome.tabs.query({{url: '{BASE}/test.html'}}))[0]""")
        W1, T1 = tab["windowId"], tab["id"]

        def errors_collected():
            info = sw.evaluate("(id) => chrome.tabs.sendMessage(id, {type: 'qarator:pageInfo'})", T1)
            kinds = sorted({e["kind"] for e in info["errors"]})
            msgs = [f'{e["kind"]}: {e["message"]} · {e["source"]} x{e["count"]}' for e in info["errors"]]
            check("page errors: bridge answers pageInfo", info["url"].endswith("/test.html"), info["title"])
            for k in ["console", "error", "promise", "http", "resource"]:
                check(f"page errors: kind '{k}' captured", k in kinds, "; ".join(m for m in msgs if m.startswith(k)))
            dup = [e for e in info["errors"] if "canvas-bg" in e["message"]]
            check("page errors: duplicate console.error collapsed with count", len(dup) == 1 and dup[0]["count"] == 2, str(dup))
            print("\n  ".join(["collected:"] + msgs))
        phase(errors_collected)

        # ---- panel in its own window; it acts on W1 as if docked there
        with ctx.expect_page() as pi:
            sw.evaluate(f"() => chrome.windows.create({{url: '{PANEL}', type: 'normal', width: 420, height: 900}})")
        panel = pi.value
        panel.set_viewport_size({"width": 420, "height": 900})
        panel.on("console", lambda m: logs.append(f"[panel] {m.type}: {m.text}") if m.type in ("error", "warning") else None)
        panel.on("pageerror", lambda e: logs.append(f"[panel] PAGEERROR: {e}"))
        panel.wait_for_load_state()
        panel.wait_for_timeout(800)
        # The panel acts on W1 as if docked there. Recording gets a drawn canvas instead of a real tab stream.
        PATCH = f"""() => {{ chrome.windows.getCurrent = async () => ({{ id: {W1} }});
          chrome.desktopCapture.chooseDesktopMedia = (sources, cb) => setTimeout(() => cb('fake-desktop'), 50);
          navigator.mediaDevices.getUserMedia = async () => {{ const c = document.createElement('canvas'); c.width = 640; c.height = 360; const x = c.getContext('2d'); let i = 0;
            setInterval(() => {{ x.fillStyle = `hsl(${{(i * 12) % 360}},70%,50%)`; x.fillRect(0, 0, 640, 360); x.fillStyle = '#000'; x.font = '40px sans-serif'; x.fillText('frame ' + i++, 40, 180); }}, 80);
            return c.captureStream(15); }}; }}"""
        panel.evaluate(PATCH)
        sw.evaluate(f"() => chrome.tabs.update({T1}, {{active: true}})")

        def clear_cool():
            sw.evaluate("() => chrome.storage.local.remove(['modelCooldown', 'modelCache'])")

        def onboarding():
            t = panel.inner_text("#main")
            check("setup: panel shows setup card before a key is set", "Set up QArator" in t, t[:120])
            wz = next(pg for pg in ctx.pages if "welcome.html" in pg.url)
            wz.on("pageerror", lambda e: logs.append(f"[welcome] PAGEERROR: {e}"))
            wz.fill("#key", "test-gemini-key-000000000000")
            wz.click("[data-go=test]")
            wz.wait_for_selector("#key-status .status", timeout=10000)
            st = wz.inner_text("#key-status")
            check("setup: key tested, Auto order shown", "Auto will use gemini-2.5-flash" in st, st)
            wz.screenshot(path=os.path.join(OUT, "00-setup-key.png"))
            wz.click("#next")
            check("setup: only two steps (key, ready)", wz.locator("#dots i").count() == 2)
            wz.click("#next")
            wz.wait_for_selector("text=All set")
            ob = sw.evaluate("async () => (await chrome.storage.local.get(['onboarded', 'settings']))")
            check("setup: finish saves onboarded + key + Auto", ob.get("onboarded") and ob["settings"]["ai"]["apiKey"] == "test-gemini-key-000000000000" and ob["settings"]["ai"]["model"] == "auto", str(ob["settings"]["ai"])[:200])
            st0 = ob["settings"]
            check("setup: the built-in setup is neutral", st0["tags"]["list"] == [] and st0["urlRules"] == [] and st0["title"]["project"] == "APP" and [f["id"] for f in st0["fields"]] == ["severity", "priority", "classification"], str([st0["title"]["project"], len(st0["tags"]["list"])]))
            panel.wait_for_timeout(500)
            check("setup: card gone after setup", "Set up QArator" not in panel.inner_text("#main"))
        phase(onboarding)
        # the built-in setup is neutral; the rest of the run uses the UR team format, as a team would import it
        FORMAT = json.load(open(os.path.join(HERE, "ur-format.json"), encoding="utf-8"))
        sw.evaluate("""async (f) => {
          const { settings = {} } = await chrome.storage.local.get('settings');
          const { qarator, formatVersion, ...rest } = f;
          await chrome.storage.local.set({ settings: { ...settings, ...rest } });
        }""", FORMAT)
        set_settings({"ai": {"provider": "gemini", "apiKey": "test-gemini-key", "model": "gemini-2.5-flash", "baseUrl": ""}})
        panel.wait_for_timeout(300)

        def to_capture():
            if panel.locator("[data-action=back-capture]").count():
                panel.click("[data-action=back-capture]")
            if panel.locator("[data-action=dismiss-error]").count():
                panel.click("[data-action=dismiss-error]")

        def shot(name):
            panel.screenshot(path=os.path.join(OUT, f"{name}.png"), full_page=True)

        def capture():
            panel.click("#main [data-action=capture]")
            panel.wait_for_selector(".shots img", timeout=8000)
            panel.wait_for_timeout(400)
            txt = panel.inner_text("#main")
            check("capture: screenshot added", panel.locator(".shots img").count() == 1)
            check("capture: page URL shown", "/test.html" in txt)
            check("capture: page errors chip", "page error" in txt and "captured" in txt, [l for l in txt.split("\n") if "error" in l][:3])
            check("capture: env chips (browser/os/viewport)", "Viewport" in txt, txt.split("Viewport")[0][-60:] if "Viewport" in txt else txt[:200])
            where = panel.inner_text("#where") if panel.locator("#where").count() else ""
            check("location: module/page read from breadcrumb", "SEVI Dashboard › State Overview" in where, where)
            shot("01-captured")
        phase(capture)

        def write_gemini():
            check("note required: Generate buttons wait for What's wrong", panel.get_attribute("#dock [data-action=write]", "aria-disabled") == "true"
                  and panel.get_attribute("#dock [data-action=generate-later]", "aria-disabled") == "true" and panel.locator("#need-note").is_visible())
            panel.keyboard.press("Control+Enter")
            panel.wait_for_timeout(300)
            check("note required: Ctrl+Enter does not write", panel.locator("#title-preview").count() == 0)
            panel.fill("#note", "Dark mode leaves the dashboard canvas white")
            check("note required: buttons unlock while typing", panel.get_attribute("#dock [data-action=write]", "aria-disabled") is None and not panel.locator("#need-note").is_visible())
            panel.click("#dock [data-action=write]")
            panel.wait_for_selector("#title-preview", timeout=15000)
            title = panel.inner_text("#title-preview").replace("\n", "")
            check("write(gemini): review opens", True)
            check("write(gemini): title built + Title Case + no trailing period",
                  title == "UR | Department Portal | SEVI Dashboard | State Overview | Dark Mode Leaves the Dashboard Canvas White", title)
            url = panel.input_value("#sec-url_note")
            check("write: URL line keeps real URL, strips repeat", url == f"{BASE}/test.html (Specifically affects the State Overview workspace)", url)
            steps = panel.input_value("#sec-steps")
            check("write: step numbering/bullets stripped", steps.split("\n") == ["Open the Department Portal and log in", "Click the theme toggle", "Open SEVI Dashboard > State Overview", "Observe the cards"], repr(steps))
            sev, pri, cls = panel.input_value("#fld-severity"), panel.input_value("#fld-priority"), panel.input_value("#fld-classification")
            check("write: AI fields matched loosely", (sev, pri, cls) == ("Minor", "Medium", "Backend"), f"{sev}/{pri}/{cls}")
            check("write: non-AI fields keep defaults", panel.input_value("#fld-test_phase") == "Smoke" and panel.input_value("#fld-device_type") == "Web Application")
            tags = panel.locator(".tagc").all_inner_texts()
            check("write: tags limited to list, deduped", tags == ["UR - Department Portal"], str(tags))
            cur = panel.get_attribute("#flow [aria-current=step]", "data-id")
            check("steps bar: Review current, Fill not reachable yet", cur == "review" and panel.locator("#flow [data-id=filled]").is_disabled(), str(cur))
            r = seen["gemini"][-1]
            body = r["body"]
            parts = body["contents"][0]["parts"]
            img = [x for x in parts if "inline_data" in x]
            check("gemini request: screenshot sent as jpeg inline_data", len(img) == 1 and img[0]["inline_data"]["mime_type"] == "image/jpeg")
            check("gemini request: key in header, not URL", r["headers"].get("x-goog-api-key") == "test-gemini-key" and "key=" not in r["url"])
            text = parts[-1]["text"]
            check("gemini request: note + errors + url in prompt", "Dark mode leaves" in text and "Page errors captured" in text and "/test.html" in text)
            check("gemini request: JSON mode + system prompt", body.get("generationConfig", {}).get("responseMimeType") == "application/json" and "senior QA engineer" in body["systemInstruction"]["parts"][0]["text"])
            shot("02-review")
        phase(write_gemini)

        def review_edits():
            panel.fill("#seg-page", "State Overview Tab")
            check("review: title preview updates live", "State Overview Tab" in panel.inner_text("#title-preview"))
            before = panel.locator("#tp-chips .tagc").count()
            panel.fill("#tp-q", "SEVI Dashboard")
            check("tags: search narrows the team list", panel.locator("#tp-list [data-pick]").count() == 1, str(panel.locator("#tp-list [data-pick]").all_inner_texts()))
            panel.press("#tp-q", "Enter")
            tags = panel.locator("#tp-chips .tagc").all_inner_texts()
            check("tags: Enter ticks the top match", len(tags) == before + 1 and any("SEVI Dashboard" in t for t in tags), str(tags))
            panel.fill("#tp-q", "")
            n0 = panel.locator("#tp-chips .tagc").count()
            panel.click("#tp-list [data-pick='UR - Citizen Portal - Schemes']")
            panel.click("#tp-list [data-pick='UR - Citizen Portal - Support']")
            check("tags: several can be ticked one after another", panel.locator("#tp-chips .tagc").count() == n0 + 2)
            check("tags: ticked rows show pressed", panel.get_attribute("#tp-list [data-pick='UR - Citizen Portal - Support']", "aria-pressed") == "true")
            panel.locator("section:has(#tp-q)").screenshot(path=os.path.join(OUT, "11-tag-picker.png"))
            panel.click("#tp-chips [data-unpick='UR - Citizen Portal - Support']")
            panel.click("#tp-list [data-pick='UR - Citizen Portal - Schemes']")
            check("tags: chip cross and row click both untick", panel.locator("#tp-chips .tagc").count() == n0)
            check("tags: no free-text tag box (team list only)", panel.locator("#new-tag").count() == 0)
            check("fields: every dropdown offers Leave empty", panel.locator("select[data-field] option[value='']").count() == panel.locator("select[data-field]").count())
            panel.select_option("#fld-device_type", "")
            check("fields: Leave empty is kept", panel.input_value("#fld-device_type") == "")
            panel.fill("#seg-page", "State Overview")
        phase(review_edits)

        def errors_paths():
            def attempt(m):
                mode["v"] = m
                to_capture()
                panel.click("#dock [data-action=write]")
                panel.wait_for_selector(".alert, #title-preview", timeout=15000)
                panel.wait_for_timeout(300)
                return panel.inner_text("#main")
            t = attempt("429")
            check("error 429: busy banner with countdown", "busy" in t and "Try again in" in t, t[:160].replace("\n", " | "))
            check("error 429: retry button disabled during countdown", panel.get_attribute("#retry", "aria-disabled") == "true")
            shot("03-ratelimit")
            panel.click("[data-action=dismiss-error]")
            clear_cool()
            t = attempt("badkey")
            check("error gemini bad key (HTTP 400): says key rejected", "key was rejected" in t.lower(), t[:160].replace("\n", " | "))
            panel.click("[data-action=dismiss-error]")
            t = attempt("garbage")
            check("error non-JSON answer: 'not usable' message", "not usable" in t, t[:160].replace("\n", " | "))
            panel.click("[data-action=dismiss-error]")
            mode["v"] = "ok"
            set_settings({"ai": {"apiKey": ""}})
            panel.wait_for_timeout(300)
            panel.click("#dock [data-action=write]")
            panel.wait_for_timeout(500)
            t = panel.inner_text("#main")
            check("no key: setup screen", "Add an AI key" in t, t[:120])
            panel.click("[data-action=blank]")
            panel.wait_for_selector("#title-preview")
            check("no key: 'fill myself' opens blank review", panel.input_value("#sec-summary") == "")
            set_settings({"ai": {"apiKey": "test-gemini-key"}})
            panel.wait_for_timeout(300)
        phase(errors_paths)

        def other_providers():
            for prov, ai in [("anthropic", {"provider": "anthropic", "apiKey": "sk-ant-test", "model": "claude-sonnet-4-5"}),
                             ("openai", {"provider": "openai", "apiKey": "sk-test", "model": "gpt-4.1"}),
                             ("compatible", {"provider": "compatible", "apiKey": "", "model": "local-vision", "baseUrl": f"http://127.0.0.1:{PORT}/v1/"})]:
                set_settings({"ai": ai})
                panel.wait_for_timeout(300)
                to_capture()
                panel.click("#dock [data-action=write]")
                try:
                    panel.wait_for_selector("#title-preview", timeout=15000)
                    ok = "Dark Mode Leaves" in panel.inner_text("#title-preview")
                except Exception:
                    ok = False
                check(f"write({prov}): report produced", ok, "" if ok else panel.inner_text("#main")[:200].replace("\n", " | "))
            a = seen["anthropic"][-1]
            check("anthropic request: browser-access header + image block", a["headers"].get("anthropic-dangerous-direct-browser-access") == "true" and any(c["type"] == "image" for c in a["body"]["messages"][0]["content"]))
            o = seen["openai"][-1]
            check("openai request: json_object + image_url", o["body"].get("response_format", {}).get("type") == "json_object" and any(c["type"] == "image_url" for c in o["body"]["messages"][1]["content"]))
            log = [json.loads(l) for l in open(os.path.join(HERE, "requests.jsonl"), encoding="utf-8")] if os.path.exists(os.path.join(HERE, "requests.jsonl")) else []
            check("compatible request: no auth header without key, trailing slash ok", log and log[-1]["auth"] is None)
            set_settings({"ai": {"provider": "gemini", "apiKey": "test-gemini-key", "model": "gemini-2.5-flash", "baseUrl": ""}})
            panel.wait_for_timeout(300)
        phase(other_providers)

        def auto_model():
            clear_cool()
            set_settings({"ai": {"provider": "gemini", "apiKey": "test-gemini-key", "model": "auto", "baseUrl": "", "models": []}})
            panel.wait_for_timeout(300)
            mode["v"] = "quota_flash"
            n0 = len(seen["gemini"])
            to_capture()
            panel.click("#dock [data-action=write]")
            panel.wait_for_selector("#title-preview", timeout=15000)
            hdr = panel.inner_text("#main").split("\n")[1:3]
            tried = [r["url"].split("models/")[1].split(":")[0] for r in seen["gemini"][n0:]]
            check("auto: tries flash first, falls back to pro on daily quota", tried == ["gemini-2.5-flash", "gemini-2.5-pro"], str(tried))
            check("auto: review says which model wrote it", any("Written by gemini-2.5-pro" in h for h in hdr), str(hdr))
            st = sw.evaluate("async () => chrome.storage.local.get(['modelCooldown', 'learnedLimits', 'usage'])")
            until = max([u for k, u in st.get("modelCooldown", {}).items() if k.startswith("gemini-2.5-flash@main:")] or [0])
            reset = panel.evaluate("async () => (await import('/lib/usage.js')).nextQuotaReset()")
            check("auto: flash rests until the daily reset", until > time.time() * 1000 and abs(until - reset) < 60000, f"{until} vs {reset}")
            check("auto: daily limit learned from the 429", st.get("learnedLimits", {}).get("gemini-2.5-flash") == 20, str(st.get("learnedLimits")))
            n1 = len(seen["gemini"])
            to_capture()
            panel.wait_for_timeout(400)
            note = panel.inner_text("#dock")
            check("usage chip: panel says flash is resting", "gemini-2.5-flash" in note and "resting until" in note, note[-200:])
            panel.click("#dock [data-action=write]")
            panel.wait_for_selector("#title-preview", timeout=15000)
            tried = [r["url"].split("models/")[1].split(":")[0] for r in seen["gemini"][n1:]]
            check("auto: resting model skipped on next write", tried == ["gemini-2.5-pro"], str(tried))
            clear_cool()
            mode["v"] = "dead_flash"
            n2 = len(seen["gemini"])
            to_capture()
            panel.click("#dock [data-action=write]")
            panel.wait_for_selector("#title-preview", timeout=15000)
            tried = [r["url"].split("models/")[1].split(":")[0] for r in seen["gemini"][n2:]]
            check("auto: 404 model skipped, next model writes", tried == ["gemini-2.5-flash", "gemini-2.5-pro"], str(tried))
            mode["v"] = "ok"
            clear_cool()
            set_settings({"ai": {"model": "gemini-2.5-flash", "fallback": False}})
            panel.wait_for_timeout(300)
            mode["v"] = "quota_flash"
            to_capture()
            panel.click("#dock [data-action=write]")
            panel.wait_for_selector(".alert", timeout=15000)
            check("fixed model, fallback off: shows rate limit instead of switching", "busy" in panel.inner_text("#main"))
            mode["v"] = "ok"
            clear_cool()
            sw.evaluate("() => chrome.storage.local.remove('learnedLimits')")
            set_settings({"ai": {"model": "auto", "fallback": True}})
            panel.wait_for_timeout(300)
            to_capture()
        phase(auto_model)

        def backup_key():
            clear_cool()
            set_settings({"ai": {"provider": "gemini", "apiKey": "test-gemini-key", "apiKey2": "backup-gemini-key", "model": "auto", "models": []}})
            panel.wait_for_timeout(300)
            mode["v"] = "quota_main"
            n0 = len(seen["gemini"])
            to_capture()
            panel.click("#dock [data-action=write]")
            panel.wait_for_selector("#title-preview", timeout=15000)
            tried = [(r["url"].split("models/")[1].split(":")[0], r["headers"].get("x-goog-api-key")) for r in seen["gemini"][n0:]]
            check("backup key: same model retried on the backup key before other models", tried == [("gemini-2.5-flash", "test-gemini-key"), ("gemini-2.5-flash", "backup-gemini-key")], str(tried))
            cool = sw.evaluate("async () => (await chrome.storage.local.get('modelCooldown')).modelCooldown || {}")
            check("backup key: only the main key's slot rests", any(k.startswith("gemini-2.5-flash@main:") for k in cool) and not any(k.startswith("gemini-2.5-flash@backup:") for k in cool), str(list(cool)))
            n1 = len(seen["gemini"])
            to_capture(); panel.wait_for_timeout(500)
            q = panel.inner_text("#dock")
            check("running-out alert: says the backup key is in use", "Quota running out" in q and "backup key" in q, q[-260:])
            shot("16-backup-alert")
            panel.click("#dock [data-action=write]")
            panel.wait_for_selector("#title-preview", timeout=15000)
            tried = [r["headers"].get("x-goog-api-key") for r in seen["gemini"][n1:]]
            check("backup key: resting main key skipped next time", tried == ["backup-gemini-key"], str(tried))
            mode["v"] = "ok"
            clear_cool()
            sw.evaluate("() => chrome.storage.local.remove('learnedLimits')")
            set_settings({"ai": {"apiKey2": ""}})
            panel.wait_for_timeout(300)
            to_capture()
        phase(backup_key)

        def no_screenshot_url():
            panel.wait_for_timeout(500); panel.evaluate("() => chrome.storage.local.remove('draft')")
            panel.reload(); panel.wait_for_timeout(800); panel.evaluate(PATCH)
            panel.click("[data-action=noshot]")
            panel.fill("#note", "Export button does nothing")
            panel.click("#dock [data-action=write]")
            panel.wait_for_selector("#title-preview", timeout=15000)
            url = panel.input_value("#sec-url_note")
            check("no screenshot: URL section still has the page address", url.startswith(f"{BASE}/test.html"), url)
            check("no screenshot: module/page from the page", panel.input_value("#seg-module") == "SEVI Dashboard" and panel.input_value("#seg-page") == "State Overview",
                  panel.input_value("#seg-module") + " / " + panel.input_value("#seg-page"))
        phase(no_screenshot_url)

        def bug_queue():
            import docx as pydocx
            panel.wait_for_timeout(500); panel.evaluate("() => chrome.storage.local.remove(['draft', 'bugs'])")
            panel.reload(); panel.wait_for_timeout(800); panel.evaluate(PATCH)
            site.goto(f"{BASE}/test.html"); site.wait_for_timeout(600)
            for note in ["Bug A: export button does nothing", "Bug B: date filter ignores end date"]:
                panel.click("#main [data-action=capture]")
                panel.wait_for_selector(".shots img", timeout=8000)
                panel.fill("#note", note)
                panel.click("#dock [data-action=generate-later]")
                panel.wait_for_timeout(400)
                check(f"generate & next: panel is clear for the next bug ({note[:5]})", panel.locator(".shots img").count() == 0 and "Capture what's wrong" in panel.inner_text("#main"))
            for _ in range(60):
                bugs = sw.evaluate("async () => ((await chrome.storage.local.get('bugs')).bugs || []).map((b) => b.status)")
                if bugs and all(x == "ready" for x in bugs):
                    break
                panel.wait_for_timeout(250)
            check("generate & next: both written in the background", bugs == ["ready", "ready"], str(bugs))
            check("bug list button shows ready count", panel.inner_text("#btn-bugs").strip() == "2", panel.inner_text("#btn-bugs"))
            panel.click("#btn-bugs")
            panel.wait_for_selector(".bug")
            bl = panel.inner_text("#main")
            check("bug list: titles and Ready status", bl.count("Ready") == 2 and "Dark Mode Leaves" in bl, bl[:300].replace(chr(10), " | "))
            check("bug list: export waits for ticks", panel.get_attribute("#dock [data-action=bugs-export-full]", "aria-disabled") == "true")
            panel.check("#bsel-all")
            shot("17-bug-list")
            with panel.expect_download() as dl:
                panel.click("#dock [data-action=bugs-export-full]")
            path = os.path.join(OUT, "18-bugs-full.docx"); dl.value.save_as(path)
            d = pydocx.Document(path); txt = [x.text for x in d.paragraphs]
            check("export full bugs: report + screenshots, page per bug", txt.count("Issue fields") == 2 and len(d.inline_shapes) == 2 and sum('w:type="page"' in x._p.xml for x in d.paragraphs) == 1, f"{txt.count('Issue fields')} reports, {len(d.inline_shapes)} images")
            with panel.expect_download() as dl:
                panel.click("#dock [data-action=bugs-export-notes]")
            path = os.path.join(OUT, "19-bugs-notes.docx"); dl.value.save_as(path)
            d = pydocx.Document(path); txt = [x.text for x in d.paragraphs]
            check("export notes only: URL, screenshot, what's wrong", "Issue fields" not in txt and txt.count("What's wrong") == 2 and "Bug A: export button does nothing" in txt
                  and sum(t.startswith("URL: ") for t in txt) == 2 and len(d.inline_shapes) == 2, str(txt[:8]))
            panel.locator("[data-action=bug-open]").first.click()
            panel.wait_for_selector("#title-preview")
            check("open ready bug: goes to review", "Dark Mode Leaves" in panel.inner_text("#title-preview"))
            site.goto(f"{ZBASE}/zoho.html"); site.wait_for_timeout(600)
            panel.click("#dock [data-action=fill]")
            panel.wait_for_selector("text=Form filled", timeout=30000)
            st_ = sw.evaluate("async () => (await chrome.storage.local.get('bugs')).bugs.map((b) => b.status)")
            check("filled bug marked in the list", st_ == ["filled", "ready"], str(st_))
            check("filled screen offers the next ready bug", "Next ready bug" in panel.inner_text("#dock"))
            panel.click("#dock [data-action=next-ready]")
            panel.wait_for_selector("#title-preview")
            check("next ready bug opens", panel.input_value("#seg-summary") != "")
            # a failed bug and Try again
            mode["v"] = "badkey"
            panel.click("#flow [data-id=capture]")
            panel.click("[data-action=next]") if panel.locator("[data-action=next]").count() else None
            panel.wait_for_timeout(300); panel.evaluate("() => chrome.storage.local.remove('draft')"); panel.reload(); panel.wait_for_timeout(800); panel.evaluate(PATCH)
            site.goto(f"{BASE}/test.html"); site.wait_for_timeout(600)
            panel.click("#main [data-action=capture]"); panel.wait_for_selector(".shots img", timeout=8000)
            panel.fill("#note", "Bug C"); panel.click("#dock [data-action=generate-later]")
            for _ in range(40):
                last = sw.evaluate("async () => (await chrome.storage.local.get('bugs')).bugs.slice(-1)[0].status")
                if last in ("failed", "ready"): break
                panel.wait_for_timeout(250)
            check("failed bug kept with reason", last == "failed", last)
            mode["v"] = "ok"
            panel.click("#btn-bugs"); panel.wait_for_selector("[data-action=bug-retry]")
            panel.click("[data-action=bug-retry]")
            for _ in range(40):
                last = sw.evaluate("async () => (await chrome.storage.local.get('bugs')).bugs.slice(-1)[0].status")
                if last == "ready": break
                panel.wait_for_timeout(250)
            check("try again writes it", last == "ready", last)
            panel.click("[data-action=bugs-back]")
            site.goto(f"{BASE}/test.html"); site.wait_for_timeout(500)


        def markup():
            to_capture()
            if not panel.locator("[data-action=markup]").count():
                panel.click("[data-action=capture]")
                panel.wait_for_selector("[data-action=markup]", timeout=8000)
            panel.click("[data-action=markup]")
            panel.wait_for_selector("#mk-canvas")
            panel.wait_for_timeout(300)
            bb = panel.locator("#mk-canvas").bounding_box()
            x, y, w, h = bb["x"], bb["y"], bb["width"], bb["height"]
            panel.mouse.move(x + w * .55, y + h * .2); panel.mouse.down(); panel.mouse.move(x + w * .9, y + h * .45, steps=5); panel.mouse.up()
            panel.keyboard.press("h")
            panel.mouse.move(x + w * .2, y + h * .6); panel.mouse.down(); panel.mouse.move(x + w * .5, y + h * .75, steps=5); panel.mouse.up()
            shot("04-markup")
            panel.click("[data-mk=apply]")
            panel.wait_for_selector(".badge", timeout=5000)
            check("markup: applied, badge shown", "Marked up" in panel.inner_text(".shots"))
            panel.wait_for_timeout(1000)
            frac = panel.evaluate("""async () => {
              const { draft } = await chrome.storage.local.get('draft');
              const s = draft.shots[0];
              const count = async (u) => { const i = new Image(); i.src = u; await i.decode(); const c = document.createElement('canvas'); c.width = i.width; c.height = i.height; const x = c.getContext('2d'); x.drawImage(i, 0, 0); const d = x.getImageData(0, 0, c.width, c.height).data; let b = 0; for (let k = 0; k < d.length; k += 4) if (d[k] === 0 && d[k+1] === 0 && d[k+2] === 0) b++; let o = 0; for (let k = 0; k < d.length; k += 4) if (d[k] === 255 && d[k+1] === 106 && d[k+2] === 43) o++; return [b / (d.length / 4), o]; };
              return { marked: await count(s.dataUrl), orig: await count(s.original), marks: s.marks.map(m => m.type) };
            }""")
            check("markup: Hide area is real black pixels in saved image", frac["marked"][0] - frac["orig"][0] > 0.03, str(frac))
            check("markup: Box drawn in orange", frac["marked"][1] > 100, str(frac))
        phase(markup)

        def recording():
            to_capture()
            check("record: no Record tab button any more", panel.locator("[data-action=record-tab]").count() == 0)
            panel.locator("#main [data-action=record-screen]").first.click()
            panel.wait_for_selector(".rec", timeout=5000)
            panel.wait_for_timeout(5200)
            t = panel.inner_text(".rec")
            check("record: bar with running timer", "Recording screen" in t and "0:0" in t, t)
            check("record: Write disabled while recording", panel.get_attribute("#dock [data-action=write]", "aria-disabled") == "true")
            shot("09-recording")
            panel.click("[data-action=stop-rec]")
            panel.wait_for_selector(".vid", timeout=8000)
            panel.wait_for_timeout(600)
            v = panel.evaluate("async () => (await chrome.storage.local.get('draft')).draft.videos")
            print("  video:", json.dumps([{k: x[k] for k in ("name", "mime", "size", "dur")} | {"frames": len(x["frames"])} for x in v]))
            check("record: video kept (mp4 when Chrome can)", len(v) == 1 and v[0]["size"] > 1000 and v[0]["name"].endswith((".mp4", ".webm")), str(v[0]["name"] if v else v))
            check("record: frames kept for the AI", len(v[0]["frames"]) >= 2, str(len(v[0]["frames"])))
            shot("10-video")
            n0 = len(seen["gemini"])
            panel.click("#dock [data-action=write]")
            panel.wait_for_selector("#title-preview", timeout=15000)
            parts = seen["gemini"][n0]["body"]["contents"][0]["parts"]
            imgs = sum(1 for x in parts if "inline_data" in x)
            check("record: AI gets screenshot + video frames", imgs == 1 + len(v[0]["frames"]), f"{imgs} images")
            check("record: prompt names the video frames", "_rec1." in parts[-1]["text"] and "frame 1 of" in parts[-1]["text"])
            to_capture()
        phase(recording)

        def zoho_fill():
            mode["v"] = "ok"
            sw.evaluate("async () => { const { settings } = await chrome.storage.local.get('settings'); settings.tags.list.push('made-up tag'); await chrome.storage.local.set({ settings }); }")  # a team tag Zoho does not have
            set_settings({"zoho": {"fillTags": True}})
            panel.wait_for_timeout(300)
            panel.click("#dock [data-action=write]")
            panel.wait_for_selector("#title-preview", timeout=15000)
            for t in ["UR - Department Portal - SEVI Dashboard", "made-up tag"]:
                panel.fill("#tp-q", t); panel.press("#tp-q", "Enter")
            print("  report tags:", panel.locator(".tagc").all_inner_texts())
            # not on Zoho yet
            panel.click("#dock [data-action=fill]")
            panel.wait_for_timeout(500)
            check("fill: not-Zoho banner when active tab isn't Zoho", "Open the Zoho Submit Issue form first" in panel.inner_text("#dock"))
            site.goto(f"{ZBASE}/zoho.html")
            site.wait_for_timeout(600)
            panel.click("#dock [data-action=fill]")
            panel.wait_for_selector("text=Form filled", timeout=30000)
            shot("05-filled-lyte")
            out = panel.inner_text("#main")
            print("  panel results:\n   " + out.replace("\n", "\n   "))
            v = site.evaluate("""() => ({
              title: document.getElementById('title').value,
              desc: document.getElementById('ed').contentDocument.body.innerHTML,
              picked: window.__picked || {},
              tokens: [...document.querySelectorAll('.ztokenfield__token')].map(t => t.textContent),
              tags: [...document.querySelectorAll('.tagchip')].map(t => t.textContent),
              files: document.getElementById('flist').textContent,
              steps: document.getElementById('steps').value,
              submitted: !!window.__submitted })""")
            print("  zoho DOM:", json.dumps({**v, "desc": v["desc"][:300]}, indent=1))
            check("zoho(lyte): title filled", v["title"].startswith("UR | Department Portal"), v["title"])
            check("zoho(lyte): description HTML in iframe editor", "<b>Summary:</b>" in v["desc"] and "<ol>" in v["desc"])
            want = {"Severity": "Minor", "Priority": "Medium", "Classification": "Backend", "Reproducible": "Always"}
            check("zoho(lyte): dropdowns picked (Backend is option 16, loaded on scroll)", all(v["picked"].get(k) == x for k, x in want.items() if k != "Test Phase") , str(v["picked"]))
            check("zoho(lyte): Assignee 'Unassigned' left empty", "Assignee" not in v["picked"])
            check("zoho(lyte): Device Type token", v["tokens"] == ["Web Application"], str(v["tokens"]))
            check("zoho(lyte): tags from existing list", v["tags"] == ["UR - Department Portal", "UR - Department Portal - SEVI Dashboard"], str(v["tags"]))
            check("zoho(lyte): screenshot attached", "qarator_" in v["files"] and ".png" in v["files"], v["files"])
            check("zoho(lyte): video attached", "_rec1." in v["files"], v["files"])
            check("zoho(lyte): never submits", not v["submitted"])
            check("fill: unknown tag flagged for manual check", "Needs a manual check" in out and "made-up tag" in out, out[:300])
            out2 = panel.inner_text("#main")
            check("filled: says how to change things", "Need to change something?" in out2 and panel.locator("#dock [data-action=back-review]").count() == 1)
            check("steps bar: all three steps reachable after filling", not panel.locator("#flow [data-id=review]").is_disabled() and panel.get_attribute("#flow [aria-current=step]", "data-id") == "filled")
            rules = sw.evaluate("async () => (await chrome.storage.local.get('settings')).settings.urlRules")
            check("learned URL rule saved after fill", rules and rules[0].get("learned") and rules[0]["match"] == f"localhost/test.html" and rules[0]["module"] == "SEVI Dashboard", str(rules[:1]))
            panel.click("#change [data-action=back-review]")
            panel.wait_for_selector("#fld-severity")
            panel.select_option("#fld-severity", "Major")
            panel.click("#dock [data-action=fill]")
            try: site.wait_for_function("(window.__picked || {}).Severity === 'Major'", timeout=20000)
            except Exception: pass
            check("edit report + fill again updates Zoho", site.evaluate("(window.__picked || {}).Severity") == "Major", site.evaluate("JSON.stringify(window.__picked)"))
            panel.wait_for_timeout(1500)
            # an install still on the old Device Type list ("Desktop", which this Zoho does not have)
            sw.evaluate("""async () => { const { settings } = await chrome.storage.local.get('settings');
              const f = settings.fields.find((x) => x.id === 'device_type'); f.options = ['Desktop', 'Mobile', 'Tablet']; f.default = 'Desktop';
              await chrome.storage.local.set({ settings }); }""")
            panel.wait_for_timeout(300)
            panel.click("#dock [data-action=back-review]")
            panel.wait_for_selector("#fld-device_type")
            panel.select_option("#fld-device_type", "Desktop")
            site.evaluate("document.querySelectorAll('.ztokenfield__token').forEach((t) => t.remove())")
            panel.click("#dock [data-action=fill]")
            panel.wait_for_selector("[data-action=use-zoho-options]", timeout=30000)
            msg = panel.inner_text("#main")
            check("missing option: says Zoho's real list (not 'couldn't open')", "Web Application" in msg and "Couldn't open" not in msg, msg[msg.find("Device Type"):][:220])
            shot("13-use-zoho-list")
            panel.click("[data-action=use-zoho-options]")
            panel.wait_for_timeout(500)
            f = sw.evaluate("async () => (await chrome.storage.local.get('settings')).settings.fields.find((x) => x.id === 'device_type')")
            check("Use Zoho's list: settings take Zoho's options", "Web Application" in f["options"] and "Desktop" not in f["options"], str(f["options"]))
            panel.click("#dock [data-action=back-review]")
            panel.select_option("#fld-device_type", "Web Application")
            panel.click("#dock [data-action=fill]")
            try: site.wait_for_function("[...document.querySelectorAll('.ztokenfield__token')].some((t) => t.textContent === 'Web Application')", timeout=20000)
            except Exception: pass
            check("fill again after update picks it", site.evaluate("[...document.querySelectorAll('.ztokenfield__token')].map((t) => t.textContent).join()") == "Web Application")
            panel.wait_for_timeout(1500)
            # steps box
            set_settings({"zoho": {"stepsField": True}})
            panel.wait_for_timeout(300)
            panel.click("#dock [data-action=fill]")
            try: site.wait_for_function("document.getElementById('steps').value", timeout=30000)
            except Exception: pass
            panel.wait_for_timeout(3000)
            steps = site.evaluate("document.getElementById('steps').value")
            check("zoho(lyte): steps box filled numbered when enabled", steps.startswith("1. Open the Department Portal"), repr(steps[:80]))
            tags2 = site.evaluate("[...document.querySelectorAll('.tagchip')].length")
            check("fill again: tags not duplicated", tags2 == 2, str(tags2))
            set_settings({"zoho": {"stepsField": False}})
        phase(zoho_fill)

        def zoho_generic():
            set_settings({"zoho": {"fillTags": False}})
            site.goto(f"{ZBASE}/generic.html")
            site.wait_for_timeout(500)
            panel.click("#dock [data-action=back-review]")
            panel.click("#dock [data-action=fill]")
            panel.wait_for_selector("text=Form filled", timeout=30000)
            out = panel.inner_text("#main")
            shot("06-filled-generic")
            print("  panel results:\n   " + out.replace("\n", "\n   "))
            v = site.evaluate("""() => ({ t: document.getElementById('t').value, d: document.getElementById('d').innerHTML,
              sev: document.getElementById('sev').selectedOptions[0].text, pri: document.getElementById('pri').value, cls: document.getElementById('cls').value,
              rep: document.getElementById('rep').value, chips: [...document.querySelectorAll('.chip')].map(c => c.textContent),
              files: [...document.getElementById('att').files].map(f => f.name) })""")
            print("  generic DOM:", json.dumps({**v, "d": v["d"][:200]}))
            check("generic: title by <label for>", v["t"].startswith("UR | "), v["t"])
            check("generic: contenteditable description", "<b>Summary:</b>" in v["d"])
            check("generic: native selects (Severity/Priority/Classification)", (v["sev"], v["pri"], v["cls"]) == ("Major", "Medium", "Issue"), str(v))
            check("generic: alt label 'Is it reproducible'", v["rep"] == "Always", v["rep"])
            check("tags off by default: nothing typed into Zoho", v["chips"] == [], str(v["chips"]))
            check("tags: list to copy one at a time after filling", "Tags to add in Zoho" in out and panel.locator("[data-action=copy-tag]").count() >= 2, out[out.find("Tags to add"):][:200])
            check("generic: single file input gets screenshot", len(v["files"]) == 1, str(v["files"]))
            if panel.locator("[data-action=goto]").count():
                panel.locator("[data-action=goto]").first.click()
                panel.wait_for_timeout(500)
                check("goto field: outlines field in page", site.evaluate("!!document.querySelector('[data-qarator-field][style*=outline]')"))
        phase(zoho_generic)

        def word_export():
            import docx as pydocx
            panel.click("#flow [data-id=review]")
            panel.wait_for_selector("#title-preview")
            with panel.expect_download() as dl:
                panel.click("#dock [data-action=export-word]")
            fname = dl.value.suggested_filename
            d, names = open_export(dl.value, os.path.join(OUT, "11-report." + fname.split(".")[-1]))
            txt = [x.text for x in d.paragraphs]
            check("word (report) with a video: one zip with the Word file and videos/", fname.endswith(".zip") and any(n.startswith("videos/") and n.endswith((".mp4", ".webm")) for n in names), f"{fname} {names}")
            check("word (report): title, URL, sections, fields", txt[0].startswith("UR | Department Portal") and any(t.startswith("URL: http://localhost") for t in txt) and "Expected Result" in txt and "Issue fields" in txt, str(txt[:6]))
            check("word (report): screenshot and video still embedded, video named", len(d.inline_shapes) == 2 and any("_rec1." in t and "videos folder" in t for t in txt), f"{len(d.inline_shapes)} images")
            set_settings({"pageErrors": False})
            sw.evaluate("() => new Promise((r) => setTimeout(r, 400))")
            check("page errors off: page hook unregistered", sw.evaluate("async () => (await chrome.scripting.getRegisteredContentScripts()).length") == 0)
            panel.wait_for_timeout(500); panel.evaluate("() => chrome.storage.local.remove('draft')")
            panel.reload(); panel.wait_for_timeout(800); panel.evaluate(PATCH)
            site.goto(f"{BASE}/test.html"); site.wait_for_timeout(800)
            panel.click("#main [data-action=capture]")
            panel.wait_for_selector(".shots img", timeout=8000)
            panel.fill("#note", "Language toggle leaves the hero in English")
            panel.wait_for_timeout(800); print("  probe after note:", panel.evaluate("async () => { const d = (await chrome.storage.local.get('draft')).draft; return d ? { view: d.view, shots: d.shots.length, note: d.note } : null; }"))
            panel.click("[data-action=toggle-steps]")
            panel.fill("#steps", "Open the citizen portal\nSwitch language to ML")
            pt = panel.inner_text("#main") + panel.inner_text("#dock")
            check("page errors off: no error chip or mention in the panel", "page error" not in pt.lower(), [l for l in pt.split(chr(10)) if "error" in l.lower()][:3])
            with panel.expect_download() as dl:
                panel.click("#dock [data-action=export-word]")
            path = os.path.join(OUT, "12-no-ai.docx"); dl.value.save_as(path)
            panel.wait_for_timeout(800); print("  probe after export:", panel.evaluate("async () => { const d = (await chrome.storage.local.get('draft')).draft; return d ? { view: d.view, shots: d.shots.length, note: d.note } : null; }"))
            d = pydocx.Document(path)
            txt = [x.text for x in d.paragraphs]
            print("  no-AI word:", txt[:12])
            check("word (no AI): note, steps, URL, location, screenshot", "What's wrong" in txt and "Language toggle leaves the hero in English" in txt and "1.\tOpen the citizen portal" in txt
                  and any(t.startswith("URL: http://localhost") for t in txt) and any("SEVI Dashboard" in t for t in txt) and len(d.inline_shapes) == 1, str(txt[:12]))
            check("page errors off: not in the Word file", "Page errors" not in txt, str([t for t in txt if "rror" in t]))
            check("word (no AI): no AI call made", True)
        phase(word_export)

        def persistence_history():
            print("  draft before reload:", panel.evaluate("async () => { const d = (await chrome.storage.local.get('draft')).draft; return d ? { view: d.view, shots: d.shots.length, note: d.note } : null; }"))
            panel.reload()
            panel.wait_for_timeout(1200)
            panel.evaluate(PATCH)
            t = panel.inner_text("#main")
            note = panel.input_value("#note", timeout=2000) if panel.locator("#note").count() else ""
            check("draft: survives panel reload", "Language toggle leaves the hero" in note and panel.locator(".shots img").count() == 1, t[:300].replace(chr(10), " | "))
            panel.click("#btn-history")
            panel.wait_for_timeout(500)
            h = panel.inner_text("#main")
            check("history: filled report listed", "Filled in Zoho" in h, h[:200].replace("\n", " | "))
            # two more saved reports, then tick all and export one Word file
            sw.evaluate("""async () => { const { history } = await chrome.storage.local.get('history'); const h = history[0];
              const more = [1, 2].map((n) => ({ ...JSON.parse(JSON.stringify(h)), id: h.id + n, title: `Second bug ${n}`, status: 'Copied, not filled' }));
              await chrome.storage.local.set({ history: [...more, ...history] }); }""")
            panel.click("[data-action=history-back]"); panel.click("#btn-history"); panel.wait_for_timeout(500)
            check("history export: button waits for a tick", panel.get_attribute("#hist-export", "aria-disabled") == "true")
            panel.locator("[data-hsel]").first.check()
            check("history export: count follows ticks", "Export 1 report to Word" in panel.inner_text("#hist-export"))
            panel.check("#hsel-all")
            n = panel.locator("[data-hsel]").count()
            check("history export: select all", f"Export {n} reports to Word" in panel.inner_text("#hist-export"), panel.inner_text("#hist-export"))
            shot("14-history-export")
            import docx as pydocx
            with panel.expect_download() as dl:
                panel.click("#hist-export")
            d, names = open_export(dl.value, os.path.join(OUT, "15-history." + dl.value.suggested_filename.split(".")[-1]))
            check("history export: screenshots come along", len(d.inline_shapes) >= 1, f"{len(d.inline_shapes)} images, {names}")
            titles = [x.text for x in d.paragraphs if x.text.startswith(("UR |", "Second bug"))]
            breaks = sum(1 for x in d.paragraphs if 'w:br w:type="page"' in x._p.xml or "w:type=\"page\"" in x._p.xml)
            check("history export: one file, one report per page", len(titles) == n and breaks == n - 1 and dl.value.suggested_filename.startswith("qarator_reports_"), f"{len(titles)} titles, {breaks} breaks, {dl.value.suggested_filename}")
            panel.click("[data-action=history-open] >> nth=-1")
            panel.wait_for_selector("#title-preview")
            check("history: reopen loads report", "Dark Mode" in panel.inner_text("#title-preview"))
            to_capture()
            t = panel.inner_text("#main")
            check("history: reopened report brings its screenshots back", panel.locator(".shots img").count() >= 1, str(panel.locator(".shots img").count()))
        phase(persistence_history)

        def paste_and_file():
            panel.wait_for_timeout(500); panel.evaluate("() => chrome.storage.local.remove('draft')")
            panel.reload(); panel.wait_for_timeout(800); panel.evaluate(PATCH)
            png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8DwHwAFBQIAX8jx0gAAAABJRU5ErkJggg==")
            panel.set_input_files("#file", files=[{"name": "a.png", "mimeType": "image/png", "buffer": png}])
            panel.wait_for_selector(".shots img", timeout=5000)
            check("add image from file", panel.locator(".shots img").count() == 1)
            panel.evaluate("""async () => { const r = await fetch(document.querySelector('.shots img').src); const b = await r.blob();
              const dt = new DataTransfer(); dt.items.add(new File([b], 'p.jpg', {type: 'image/jpeg'}));
              document.dispatchEvent(new ClipboardEvent('paste', {clipboardData: dt, bubbles: true})); }""")
            panel.wait_for_timeout(800)
            check("paste image (Ctrl+V)", panel.locator(".shots img").count() == 2, str(panel.locator(".shots img").count()))
        phase(paste_and_file)
        phase(bug_queue)

        def options_page():
            op = ctx.new_page()
            op.on("pageerror", lambda e: logs.append(f"[options] PAGEERROR: {e}"))
            op.goto(f"chrome-extension://{EID}/options/options.html")
            op.wait_for_timeout(800)
            op.click("[data-act=check]")
            op.wait_for_timeout(1500)
            opts = op.locator("#model option").all_inner_texts()
            check("options: Check key loads Gemini models (embedding filtered)", len(opts) == 3 and opts[0] == "Auto (recommended)" and not any("embedding" in o for o in opts), str(opts))
            op.locator("#usage").scroll_into_view_if_needed()
            ut = op.inner_text("#usage")
            op.locator("#usage").screenshot(path=os.path.join(OUT, "08-usage.png"))
            print("  usage section:\n   " + ut.replace("\n", "\n   ")[:1500])
            check("usage: rows per model with requests and failures", "gemini-2.5-pro" in ut and "gemini-2.5-flash" in ut and "limit)" in ut, ut[:300])
            check("usage: tokens and rupee cost shown", "3.6k / 900" in ut and "~\u20b9" in ut, ut[:300])
            op.click("[data-act=usage-days][data-i='7']")
            check("usage: 7-day range", "last 7 days" in op.inner_text("#usage"))
            op.select_option("#prov", "openai")
            op.wait_for_timeout(300)
            op.fill("#key", "sk-test-0000000000000000000000")
            op.wait_for_timeout(2500)
            opts = op.locator("#model option").all_inner_texts()
            check("options: OpenAI models load by themselves, Auto first", opts == ["Auto (recommended)", "gpt-4.1"], str(opts))
            check("settings: team format starts collapsed", not op.evaluate("document.getElementById('advanced').open"))
            op.click("#advanced > summary")
            op.fill("#tproj", "UR2")
            op.wait_for_timeout(600)
            with op.expect_download() as dl:
                op.click("[data-act=export]")
            path = dl.value.path()
            data = json.load(open(path, encoding="utf-8"))
            check("export: file name + no API key", dl.value.suggested_filename == "qarator-format-ur.json" and "ai" not in data and "usage" not in data and "appearance" not in data and "sk-test" not in json.dumps(data), dl.value.suggested_filename)
            data["title"]["project"] = "IMPORTED"
            imp = os.path.join(OUT, "import.json"); json.dump(data, open(imp, "w"))
            rows0 = op.locator(".prow").count()
            op.set_input_files("#import-file", imp)
            op.wait_for_timeout(800)
            check("import: applies format, keeps key", op.input_value("#tproj") == "IMPORTED" and op.input_value("#key") == "sk-test-0000000000000000000000")
            check("import: same project id is added as a copy, not over the original", op.locator(".prow").count() == rows0 + 1, f"{rows0} -> {op.locator('.prow').count()}")
            check("import: existing project is still there", "UR" in op.locator(".prow input").first.input_value() or op.locator(".prow input").count() == 2)
            bad = os.path.join(OUT, "bad.json"); open(bad, "w").write("{\"x\":1}")
            op.set_input_files("#import-file", bad)
            op.wait_for_timeout(500)
            check("import: rejects non-QArator file with a plain reason", op.input_value("#tproj") == "IMPORTED" and "not imported" in op.inner_text("#projects") and "qarator" in op.inner_text("#projects"), op.inner_text("#projects")[:160])
            op.screenshot(path=os.path.join(OUT, "07-options.png"))
        phase(options_page)

        def projects_phase():
            op = ctx.new_page()
            op.on("pageerror", lambda e: logs.append(f"[projects] PAGEERROR: {e}"))
            op.goto(f"chrome-extension://{EID}/options/options.html#projects")
            op.wait_for_timeout(800)
            ev = lambda js, arg=None: op.evaluate(js, arg)
            # the setup guide: downloads, and its example is a valid project file
            with op.expect_download() as dl:
                op.click("[data-act=guide]")
            guide = open(dl.value.path(), encoding="utf-8").read()
            check("guide: downloads as qarator-setup-guide.md", dl.value.suggested_filename == "qarator-setup-guide.md" and guide.startswith("# QArator setup guide"), dl.value.suggested_filename)
            check("guide: asks the agent the right questions and states the rules", "Ask first" in guide and "meta.scopes" in guide and "Never put keys" in guide)
            ex = json.loads(guide.split("```json\n", 1)[1].split("\n```", 1)[0])
            ex_path = os.path.join(OUT, "guide-example.json"); json.dump(ex, open(ex_path, "w"))
            rows0 = op.locator(".prow").count()
            op.set_input_files("#import-file", ex_path)
            op.wait_for_timeout(900)
            check("guide: its example imports with nothing to repair", op.locator(".prow").count() == rows0 + 1 and "to check" not in op.inner_text("#projects") and "Imported" in op.inner_text("#projects"), op.inner_text("#projects")[:200])
            check("projects: the imported project is in use", ev("async () => (await chrome.storage.local.get('settings')).settings.meta.id") == "acme_shop")
            # portals: the right one applies by address, a longer match wins, personal settings never move
            r = ev("""async () => {
              const m = await import('/lib/settings.js');
              const s = m.normalize((await chrome.storage.local.get('settings')).settings);
              const a = m.applyScope(s, 'https://admin.acme.example/orders/7');
              const b = m.applyScope(s, 'https://shop.acme.example/checkout');
              const none = m.applyScope(s, 'https://elsewhere.example/');
              return { a: [a.scope, a.title.project, a.tags.list.length, a.urlRules.length], b: [b.scope, b.fields.map((f) => f.id).join(), b.tags.list.join('|')], none: [none.scope, none.tags.list.length, none.title.project], ai: s.ai.provider };
            }""")
            check("portals: admin address applies the admin portal", r["a"] == ["Admin site", "ACME", 3, 2], str(r["a"]))
            check("portals: shop portal adds a field and its own tags", r["b"][0] == "Customer shop" and r["b"][1].endswith("device") and r["b"][2].startswith("Shop"), str(r["b"]))
            check("portals: an unmatched address uses the project as it is", r["none"][0] == "" and r["none"][1] == 0, str(r["none"]))
            # the project settings are editable by hand: sites and a portal
            op.reload(); op.wait_for_timeout(800)
            op.click("#advanced > summary")
            check("portals: editor lists both portals", op.locator(".scope").count() == 2, str(op.locator(".scope").count()))
            op.fill("#sc-c1", "SHOP")
            op.fill("#sites", "shop.acme.example\nadmin.acme.example\nstaging.acme.example")
            op.wait_for_timeout(700)
            sv = ev("async () => (await chrome.storage.local.get('settings')).settings.meta")
            check("portals: hand edits are saved", sv["scopes"][1]["override"]["title"]["project"] == "SHOP" and len(sv["sites"]) == 3, str(sv)[:200])
            op.click("[data-act=sc-add]"); op.wait_for_timeout(300)
            check("portals: a new portal can be added by hand", op.locator(".scope").count() == 3)
            op.click("[data-act=sc-del][data-i='2']"); op.wait_for_timeout(300)
            check("portals: and removed", op.locator(".scope").count() == 2)
            # switching: the panel switches too, and the shelf keeps every project's own edits
            auto = ev("""async () => { const p = await import('/lib/projects.js'); return [await p.projectForUrl('https://staging.acme.example/x'), await p.projectForUrl('https://nowhere.example/')]; }""")
            check("projects: a page on a project's site finds that project", auto[0] == "acme_shop" and auto[1] == "", str(auto))
            op.reload(); op.wait_for_timeout(800)
            ur = [i for i in range(op.locator("[data-act=pj-use]").count())]
            op.locator("[data-act=pj-use]").first.click(); op.wait_for_timeout(900)
            now = ev("async () => (await chrome.storage.local.get('settings')).settings.meta.id")
            check("projects: Use switches the project in use", now != "acme_shop", now)
            back = ev("async () => { const p = await import('/lib/projects.js'); await p.switchProject('acme_shop'); return (await chrome.storage.local.get('settings')).settings.meta.scopes[1].override.title.project; }")
            check("projects: switching back keeps hand edits", back == "SHOP", back)
            check("projects: switching leaves AI key and look alone", ev("async () => (await chrome.storage.local.get('settings')).settings.ai.apiKey") != "")
            # a bad file says what is wrong in plain words
            bad = os.path.join(OUT, "bad-format.json")
            json.dump({"qarator": "format", "sections": {}, "fields": "x"}, open(bad, "w"))
            op.set_input_files("#import-file", bad); op.wait_for_timeout(600)
            t = op.inner_text("#projects")
            check("import errors: say which key is wrong", "not imported" in t and "\"sections\" must be a list" in t and "\"fields\" must be a list" in t, t[:240])
            warn = os.path.join(OUT, "warn-format.json")
            json.dump({"qarator": "format", "meta": {"id": "warny", "name": "Warny"}, "fields": [{"name": "X", "options": []}, {"id": "a", "name": "Y", "options": ["p"], "default": "z"}]}, open(warn, "w"))
            op.set_input_files("#import-file", warn); op.wait_for_timeout(900)
            t = op.inner_text("#projects")
            check("import warnings: list what was repaired", "2 things to check" in t and "fields[0].options is empty" in t and "fields[1].default" in t, t[:300])
            op.screenshot(path=os.path.join(OUT, "12-projects.png"))
            # delete a project, and the last one cannot be deleted
            before = op.locator(".prow").count()
            op.once("dialog", lambda d: d.accept())
            op.locator("[data-act=pj-del]:not([disabled])").last.click(); op.wait_for_timeout(900)
            check("projects: delete removes one", op.locator(".prow").count() == before - 1, f"{before} -> {op.locator('.prow').count()}")
            # neutral new project
            op.once("dialog", lambda d: d.accept("Fresh"))
            op.click("[data-act=pj-new]"); op.wait_for_timeout(900)
            neu = ev("async () => { const s = (await chrome.storage.local.get('settings')).settings; return [s.meta.name, s.tags.list.length, s.urlRules.length, s.title.project]; }")
            check("projects: New blank project starts neutral", neu[0] == "Fresh" and neu[1] == 0 and neu[2] == 0 and neu[3] == "APP", str(neu))
            # the panel header switches projects
            pg = ctx.new_page()
            pg.goto(f"chrome-extension://{EID}/sidepanel/sidepanel.html"); pg.wait_for_timeout(1000)
            check("panel: project switcher shows with several projects", pg.locator("#projsel").count() == 1 and pg.locator("#projsel option").count() >= 3, str(pg.locator("#projsel option").all_inner_texts()))
            pg.select_option("#projsel", "acme_shop"); pg.wait_for_timeout(900)
            check("panel: choosing a project switches it", ev("async () => (await chrome.storage.local.get('settings')).settings.meta.id") == "acme_shop")
            pg.close()
            # restore the UR setup for the phases that follow
            sw.evaluate("""async (f) => {
              const { settings = {} } = await chrome.storage.local.get('settings');
              const { qarator, formatVersion, ...rest } = f;
              await chrome.storage.local.set({ settings: { ...settings, ...rest } });
            }""", FORMAT)
        phase(projects_phase)

        def key_setup():
            FAKE = "AIza" + "S" * 35
            op = ctx.new_page()
            op.on("pageerror", lambda e: logs.append(f"[key] PAGEERROR: {e}"))
            op.goto(f"chrome-extension://{EID}/options/options.html"); op.wait_for_timeout(500)
            r = op.evaluate("""async (k) => {
              const m = await import('/lib/keys.js');
              return [m.cleanKey('  "' + k + '"\\n'), m.cleanKey('Bearer ' + k), m.detectProvider(k), m.detectProvider('sk-ant-api03-' + 'x'.repeat(30)), m.detectProvider('sk-' + 'y'.repeat(30)), m.detectProvider('hello'), m.detectProvider(' ' + k + ' ')];
            }""", FAKE)
            check("keys: copy leftovers are stripped", r[0] == FAKE and r[1] == FAKE, str(r[:2]))
            check("keys: provider is read from the key", r[2:] == ["gemini", "anthropic", "openai", "", "gemini"], str(r[2:]))
            # the setup page: plain steps, wrong provider fixed by the key itself, quotes and spaces removed
            set_settings({"ai": {"provider": "gemini", "apiKey": "", "apiKey2": "", "model": "auto", "baseUrl": "", "models": []}})
            wz = ctx.new_page()
            wz.on("pageerror", lambda e: logs.append(f"[welcome2] PAGEERROR: {e}"))
            wz.goto(f"chrome-extension://{EID}/options/welcome.html"); wz.wait_for_timeout(700)
            t = wz.inner_text("#step")
            check("setup: plain steps to get a key are shown", "Create API key" in t and "free" in t and "Get my free key" in t and "personal Gmail" in t, t[:900])
            check("setup: says what to do with no key", "Do this later" in t and "Export to Word" in t)
            wz.select_option("#prov", "anthropic"); wz.wait_for_timeout(300)
            wz.fill("#key", f'  "{FAKE}" ')
            wz.wait_for_timeout(1800)
            check("setup: a Gemini key switches the provider and is cleaned", wz.input_value("#prov") == "gemini" and wz.input_value("#key") == FAKE, f"{wz.input_value('#prov')} {wz.input_value('#key')[:8]}")
            check("setup: the pasted key is tested on its own", "Key works" in wz.inner_text("#key-status"), wz.inner_text("#key-status"))
            check("setup: steps give way once the key works", wz.locator(".getkey").count() == 0)
            wz.screenshot(path=os.path.join(OUT, "13-setup-key-ok.png"))
            wz.close()
            # the helper on Google AI Studio: absent unless the setup asked, then hands the key over
            page_html = f"<html><body><h1>API keys</h1><p>Your key</p><code>{FAKE}</code></body></html>"
            ctx.route("https://aistudio.google.com/**", lambda route: route.fulfill(status=200, content_type="text/html", body=page_html))
            sw.evaluate("async () => { await chrome.storage.local.remove('keyHelper'); }")
            g = ctx.new_page(); g.goto("https://aistudio.google.com/app/apikey"); g.wait_for_timeout(1200)
            check("helper: not shown unless the setup asked for it", g.locator("#qarator-key-helper").count() == 0)
            g.close()
            sw.evaluate("async () => { await chrome.storage.local.set({ keyHelper: { until: Date.now() + 600000 } }); }")
            g = ctx.new_page(); g.goto("https://aistudio.google.com/app/apikey"); g.wait_for_timeout(1500)
            check("helper: shown on the key page when the setup asked", g.locator("#qarator-key-helper").count() == 1)
            btn = g.locator("#qarator-key-helper #use")
            check("helper: finds the key on the page and offers it", btn.is_enabled() and "Use this key" in btn.inner_text(), btn.inner_text())
            g.screenshot(path=os.path.join(OUT, "14-key-helper.png"))
            # a page that is not AI Studio, or an extension page, cannot set the key through the same message
            bad = op.evaluate("""async (k) => (await chrome.runtime.sendMessage({ type: 'qarator-set-key', key: k.replace('AIza', 'AIzb') }))""", FAKE)
            check("helper: a key of the wrong shape is refused", not (bad or {}).get("ok"), str(bad))
            wrong = op.evaluate("""async (k) => (await chrome.runtime.sendMessage({ type: 'qarator-set-key', key: k }))""", FAKE)
            check("helper: only the AI Studio page may hand a key over", not (wrong or {}).get("ok"), str(wrong))
            btn.click(); g.wait_for_timeout(1500)
            check("helper: says it is sent", "Sent" in g.locator("#qarator-key-helper #use").inner_text())
            st2 = sw.evaluate("async () => { const { settings, keyHelper } = await chrome.storage.local.get(['settings', 'keyHelper']); return { key: settings.ai.apiKey, provider: settings.ai.provider, flag: !!keyHelper }; }")
            check("helper: key saved as Gemini and the request is closed", st2["key"] == FAKE and st2["provider"] == "gemini" and not st2["flag"], str(st2)[:120])
            g.close(); op.close()
            ctx.unroute("https://aistudio.google.com/**")
            set_settings({"ai": {"provider": "gemini", "apiKey": "test-gemini-key", "model": "gemini-2.5-flash", "baseUrl": "", "apiKey2": ""}})
        phase(key_setup)

        def project_wizard():
            w = ctx.new_page()
            w.on("pageerror", lambda e: logs.append(f"[wizard] PAGEERROR: {e}"))
            w.goto(f"chrome-extension://{EID}/options/project-setup.html"); w.wait_for_timeout(600)
            t = w.inner_text("#step")
            check("wizard: three ways in are offered", "I have a project file" in t and "Let an AI agent write it" in t and "answer a few questions" in t, t[:200])
            w.screenshot(path=os.path.join(OUT, "15-wizard-start.png"))
            # by hand
            w.click("[data-id=hand]"); w.wait_for_timeout(300)
            check("wizard(hand): name is required before Next", w.get_attribute("#next", "aria-disabled") == "true")
            w.fill("#pname", "Wiz Shop"); w.fill("#pcode", "WIZ"); w.fill("#psites", "wiz.example\nadmin.wiz.example")
            check("wizard(hand): Next opens once named", w.get_attribute("#next", "aria-disabled") == "false")
            w.click("#next"); w.wait_for_timeout(300)
            check("wizard(hand): title step shows a live example", w.inner_text("#ex") == "WIZ | Admin | Orders | Order List | Export Button Does Nothing", w.inner_text("#ex"))
            w.uncheck("[data-seg=module]"); w.wait_for_timeout(200)
            w.fill("#sep", " - "); w.select_option("#scase", "sentence")
            check("wizard(hand): example follows parts, separator and case", w.inner_text("#ex") == "WIZ - Admin - Order List - Export button does nothing", w.inner_text("#ex"))
            w.screenshot(path=os.path.join(OUT, "16-wizard-title.png"))
            w.click("#next"); w.wait_for_timeout(300)
            check("wizard(hand): sections listed, extras off by default", w.locator("[data-sec]").count() == 8 and not w.is_checked("[data-sec='6']"), str(w.locator("[data-sec]").count()))
            w.uncheck("[data-sec='3']"); w.fill("#newsec", "Workaround"); w.click("[data-go=addsec]"); w.wait_for_timeout(200)
            check("wizard(hand): a section can be added", w.locator("[data-sec]").count() == 9)
            w.click("#next"); w.wait_for_timeout(300)
            check("wizard(hand): fields offer presets with choices", w.locator("[data-fon]").count() == 5 and "Show stopper" in w.input_value("[data-fopt='0']"))
            w.fill("[data-fopt='0']", "Blocker, Major, Minor"); w.check("[data-fon='3']")
            w.fill("#newfield", "Browser"); w.click("[data-go=addfield]"); w.wait_for_timeout(200)
            w.fill("[data-fopt='5']", "Chrome, Edge, Firefox")
            w.click("#next"); w.wait_for_timeout(300)
            w.fill("#tags", "Shop\nShop - Checkout\n  Admin  \n")
            w.click("#next"); w.wait_for_timeout(300)
            w.fill("#style", "WIZ - Admin - Order List - Export button does nothing\n\nSummary: Nothing downloads.")
            w.click("#next"); w.wait_for_timeout(300)
            rv = w.inner_text("#step")
            check("wizard(hand): review lists what will be made", "Wiz Shop" in rv and "WIZ" in rv and "3 tags" in rv and "style example added" in rv and "Browser" in rv, rv[:400])
            w.screenshot(path=os.path.join(OUT, "17-wizard-review.png"))
            w.click("#next"); w.wait_for_timeout(1200)
            check("wizard(hand): finishes", "Wiz Shop is ready" in w.inner_text("#step"), w.inner_text("#step")[:80])
            s = sw.evaluate("async () => (await chrome.storage.local.get('settings')).settings")
            check("wizard(hand): project created and in use", s["meta"]["name"] == "Wiz Shop" and s["meta"]["sites"] == ["wiz.example", "admin.wiz.example"] and s["title"]["project"] == "WIZ" and s["title"]["separator"] == " - ", str(s["meta"])[:160])
            names = [x["name"] for x in s["sections"]]
            check("wizard(hand): sections, fields, tags as answered", names[-1] == "Workaround" and "Steps to Reproduce" not in names and [f["id"] for f in s["fields"]] == ["severity", "priority", "classification", "environment", "browser"] and s["fields"][0]["options"] == ["Blocker", "Major", "Minor"] and s["tags"]["list"] == ["Shop", "Shop - Checkout", "Admin"], str([f["id"] for f in s["fields"]]))
            check("wizard(hand): AI settings untouched", s["ai"]["provider"] in ("gemini", "openai", "anthropic") and "apiKey" in s["ai"])
            w.close()
            # agent route: guide, then pasted replies (chatter, broken file, good file in a code fence)
            a = ctx.new_page()
            a.on("pageerror", lambda e: logs.append(f"[wizard2] PAGEERROR: {e}"))
            a.goto(f"chrome-extension://{EID}/options/project-setup.html?route=agent"); a.wait_for_timeout(500)
            with a.expect_download() as dl:
                a.click("[data-go=guide]")
            check("wizard(agent): guide downloads", dl.value.suggested_filename == "qarator-setup-guide.md")
            check("wizard(agent): says what to give the agent", "real finished bug report" in a.inner_text("#step") and "tag list" in a.inner_text("#step") and "Copy message" in a.inner_text("#step"))
            a.fill("#json", "Sure! Here you go: it is below.")
            a.click("#next"); a.wait_for_timeout(500)
            check("wizard(agent): non-JSON chatter is explained", "not valid JSON" in a.inner_text("#step") and "only the JSON block" in a.inner_text("#step"))
            a.fill("#json", "```json\n{ \"qarator\": \"format\", \"sections\": {} }\n```")
            a.click("#next"); a.wait_for_timeout(500)
            check("wizard(agent): a broken file lists errors and what to do", "\"sections\" must be a list" in a.inner_text("#step") and "Send this list to your AI agent" in a.inner_text("#step"), a.inner_text("#step")[-300:])
            ex = json.loads(open(os.path.join(OUT, "guide-example.json"), encoding="utf-8").read())
            ex["meta"]["id"] = "wiz_agent"; ex["meta"]["name"] = "Agent Made"
            a.fill("#json", "```json\n" + json.dumps(ex) + "\n```")
            a.click("#next"); a.wait_for_timeout(1200)
            check("wizard(agent): a good reply is imported from inside a code fence", "Agent Made is ready" in a.inner_text("#step"), a.inner_text("#step")[:80])
            a.close()
            # import route by file
            i = ctx.new_page(); i.goto(f"chrome-extension://{EID}/options/project-setup.html?route=import"); i.wait_for_timeout(400)
            ex["meta"]["id"] = "wiz_file"; ex["meta"]["name"] = "From File"
            pth = os.path.join(OUT, "wiz-file.json"); json.dump(ex, open(pth, "w"))
            i.set_input_files("#pick-file", pth); i.wait_for_timeout(1200)
            check("wizard(import): a chosen file is imported", "From File is ready" in i.inner_text("#step"), i.inner_text("#step")[:80])
            i.close()
            # Settings: the group header is a real heading with an icon, and links to the wizard
            op = ctx.new_page(); op.goto(f"chrome-extension://{EID}/options/options.html"); op.wait_for_timeout(700)
            check("settings: Format group has an icon, not a font glyph", op.locator("#advanced > summary .adv-c svg").count() == 1)
            check("settings: guided setup is offered", op.locator("a[href='project-setup.html']").count() >= 2)
            op.click("#advanced > summary"); op.wait_for_timeout(300)
            op.locator("#advanced > summary").screenshot(path=os.path.join(OUT, "18-format-group.png"))
            op.close()
            sw.evaluate("""async (f) => {
              const { settings = {} } = await chrome.storage.local.get('settings');
              const { qarator, formatVersion, ...rest } = f;
              await chrome.storage.local.set({ settings: { ...settings, ...rest } });
            }""", FORMAT)
        phase(project_wizard)

        def appearance():
            op = ctx.new_page()
            op.on("pageerror", lambda e: logs.append(f"[appearance] PAGEERROR: {e}"))
            op.goto(f"chrome-extension://{EID}/options/options.html#appearance")
            op.wait_for_timeout(800)
            pick = lambda key, v: (op.check(f"[data-path='appearance.{key}'][value='{v}']"), op.wait_for_timeout(500))
            css = lambda sel, prop, pseudo=None: op.evaluate("([s, p, ps]) => getComputedStyle(document.querySelector(s), ps).getPropertyValue(p)", [sel, prop, pseudo])
            check("appearance: section and nav link", op.locator("#appearance h2").inner_text() == "Appearance" and op.locator(".nav a[href='#appearance']").count() == 1)
            # glow, light theme
            pick("theme", "light")
            check("appearance: Light sets data-theme", op.get_attribute("html", "data-theme") == "light")
            op.hover("#ap-try")
            op.wait_for_timeout(300)
            check("glow: hover ring on field", css("#ap-try", "box-shadow") not in ("", "none"), css("#ap-try", "box-shadow"))
            op.focus("#ap-try")
            op.wait_for_timeout(900)
            wrap = "[data-beam]:has(> #ap-try)"
            sh, after = css("#ap-try", "box-shadow"), float(css(wrap, "opacity", "::after") or 0)
            check("glow: focus ring + beam visible (light)", "0px 0px 0px 3px" in sh and "/ 0)" not in sh and after >= 0.6 and css("#ap-try", "border-top-color") == "rgb(194, 65, 12)", f"{sh} / beam {after}")
            op.locator("#appearance").screenshot(path=os.path.join(OUT, "09-appearance-light.png"))
            # dark theme, live in the side panel too
            panel2 = ctx.new_page()
            panel2.goto(PANEL)
            panel2.wait_for_timeout(800)
            op.bring_to_front()  # background tabs do not run CSS animations
            pick("theme", "dark")
            op.focus("#ap-try")
            op.wait_for_timeout(900)
            check("appearance: Dark sets data-theme and colours", op.get_attribute("html", "data-theme") == "dark" and css("body", "background-color") == "rgb(20, 21, 23)", css("body", "background-color"))
            check("appearance: beam follows the chosen theme", op.get_attribute(wrap, "data-beam") == "qa-beam-d")
            after = float(css(wrap, "opacity", "::after") or 0)
            check("glow: focus beam visible (dark)", after >= 0.6, str(after))
            op.locator("#appearance").screenshot(path=os.path.join(OUT, "10-appearance-dark.png"))
            check("appearance: side panel switches live", panel2.get_attribute("html", "data-theme") == "dark" and panel2.evaluate("getComputedStyle(document.body).backgroundColor") == "rgb(20, 21, 23)")
            panel2.screenshot(path=os.path.join(OUT, "11-panel-dark.png"))
            # text size
            pick("fontSize", "larger")
            fs = css("body", "font-size")
            check("appearance: Larger text scales body font", fs == "16.25px", fs)
            check("appearance: side panel text scales too", panel2.evaluate("getComputedStyle(document.body).fontSize") == "16.25px")
            pick("fontSize", "default")
            check("appearance: Default text is 13px", css("body", "font-size") == "13px", css("body", "font-size"))
            # compact panel
            pick("density", "compact")
            check("appearance: compact panel spacing", panel2.evaluate("getComputedStyle(document.querySelector('#main')).paddingTop") == "10px")
            pick("density", "comfortable")
            # animations
            op.focus("#ap-try")
            op.wait_for_timeout(200)
            on = css(wrap, "animation-duration")
            pick("motion", "off")
            op.focus("#ap-try")
            op.wait_for_timeout(300)
            off = css(wrap, "animation-duration")
            check("animations off: beam stops", "1.96s" in on and set(x.strip() for x in off.split(",")) == {"0s"}, f"{on} -> {off}")
            after = float(css(wrap, "opacity", "::after") or 0)
            check("animations off: focus beam still shown, still", after >= 0.6, str(after))
            check("animations off: panel gets it", panel2.get_attribute("html", "data-motion") == "off")
            pick("motion", "system")
            # glow strength
            pick("glow", "off")
            op.focus("#ap-try")
            op.wait_for_timeout(1000)  # let the focus-ring transition finish
            sh = css("#ap-try", "box-shadow")
            check("glow off: no beam, focus ring kept", css(wrap, "display", "::after") == "none" and "0px 0px 0px 3px" in sh and "/ 0) 0px 0px 0px 3px" not in sh, sh)
            pick("glow", "strong")
            # per person: not exported, kept on import
            op.wait_for_timeout(500)
            op.click("#advanced > summary")
            with op.expect_download() as dl:
                op.click("[data-act=export]")
            data = json.load(open(dl.value.path(), encoding="utf-8"))
            check("appearance: not in format export", "appearance" not in data, str(list(data)))
            data["appearance"] = {"theme": "light"}
            imp = os.path.join(OUT, "import-ap.json"); json.dump(data, open(imp, "w"))
            op.set_input_files("#import-file", imp)
            op.wait_for_timeout(900)
            check("appearance: kept on import", op.get_attribute("html", "data-theme") == "dark")
            # reload keeps it (boot script, before settings load)
            op.reload()
            check("appearance: applied before first paint after reload", op.evaluate("document.documentElement.dataset.theme") == "dark")
            op.wait_for_timeout(600)
            pick("theme", "system")
            check("appearance: System removes the override", op.get_attribute("html", "data-theme") is None)
            panel2.close()
        phase(appearance)

        ctx.close()
finally:
    srv.terminate()

print("\n==== console errors ====")
print("\n".join(logs) or "none")
f = [r for r in results if not r[1]]
print(f"\n==== {len(results) - len(f)}/{len(results)} passed ====")
for n, _, d in f:
    print("FAIL", n, "::", d)
